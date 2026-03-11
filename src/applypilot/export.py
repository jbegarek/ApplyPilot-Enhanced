from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from applypilot.config import APP_DIR
from applypilot.database import close_connection, get_connection


CURATED_READY_JOB_COLUMNS = [
    "title",
    "site",
    "location",
    "fit_score",
    "url",
    "application_url",
    "tailored_resume_path",
    "cover_letter_path",
    "apply_status",
    "apply_error",
    "score_reasoning",
    "discovered_at",
    "scored_at",
    "tailored_at",
    "cover_letter_at",
]


def fetch_ready_jobs_for_export(
    *,
    db_path: Path | str | None = None,
    min_score: int | None = None,
    include_failed: bool = False,
) -> list[dict[str, Any]]:
    conn = get_connection(db_path)

    clauses = [
        "tailored_resume_path IS NOT NULL",
        "applied_at IS NULL",
        "(apply_status IS NULL OR apply_status != 'applied')",
        "(apply_status IS NULL OR apply_status != 'in_progress')",
    ]
    params: list[Any] = []

    if not include_failed:
        clauses.append("(apply_status IS NULL OR apply_status != 'failed')")

    if min_score is not None:
        clauses.append("fit_score >= ?")
        params.append(min_score)

    where_sql = " AND ".join(clauses)
    rows = conn.execute(
        f"""
        SELECT *
        FROM jobs
        WHERE {where_sql}
        ORDER BY fit_score DESC, site, title, url
        """,
        params,
    ).fetchall()

    close_connection(db_path)
    return [dict(row) for row in rows]


def export_ready_jobs_to_xlsx(*, rows: list[dict[str, Any]], output: Path | str) -> Path:
    workbook = Workbook()
    curated_sheet = workbook.active
    curated_sheet.title = "ready_to_apply"
    raw_sheet = workbook.create_sheet("raw_ready_jobs")

    raw_columns = list(rows[0].keys()) if rows else CURATED_READY_JOB_COLUMNS

    _write_sheet(curated_sheet, CURATED_READY_JOB_COLUMNS, rows)
    _write_sheet(raw_sheet, raw_columns, rows)

    output_path = Path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def build_default_ready_jobs_export_path(*, base_dir: Path | str | None = None) -> Path:
    directory = Path(base_dir) if base_dir is not None else APP_DIR / "exports"
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    return directory / f"ready_jobs_{timestamp}.xlsx"


def _write_sheet(worksheet, columns: list[str], rows: list[dict[str, Any]]) -> None:
    worksheet.append(columns)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        worksheet.append([row.get(column) for column in columns])

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    for index, column in enumerate(columns, start=1):
        letter = worksheet.cell(row=1, column=index).column_letter
        width = max(len(column), 12)
        for row_index in range(2, worksheet.max_row + 1):
            value = worksheet.cell(row=row_index, column=index).value
            if value is None:
                continue
            text = str(value)
            width = min(max(width, len(text) + 2), 80)
            if column in {"url", "application_url"} and text.startswith(("http://", "https://")):
                cell = worksheet.cell(row=row_index, column=index)
                cell.hyperlink = text
                cell.style = "Hyperlink"
        worksheet.column_dimensions[letter].width = width
