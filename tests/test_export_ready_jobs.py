from __future__ import annotations

from pathlib import Path

import applypilot.cli as cli
import applypilot.database as database
from openpyxl import load_workbook
from typer.testing import CliRunner


runner = CliRunner()


def test_export_ready_jobs_command_exists(monkeypatch, tmp_path) -> None:
    monkeypatch.setattr(cli, "_bootstrap", lambda: None)
    monkeypatch.setattr(cli, "_export_ready_jobs", lambda **_kwargs: None, raising=False)

    result = runner.invoke(
        cli.app,
        ["export", "ready-jobs", "--output", str(tmp_path / "ready.xlsx")],
    )

    assert result.exit_code == 0


def test_export_ready_jobs_calls_export_helper(monkeypatch, tmp_path) -> None:
    called: dict = {}

    monkeypatch.setattr(cli, "_bootstrap", lambda: None)
    monkeypatch.setattr(
        cli,
        "_export_ready_jobs",
        lambda **kwargs: called.update(kwargs),
        raising=False,
    )

    output = tmp_path / "ready.xlsx"
    result = runner.invoke(
        cli.app,
        ["export", "ready-jobs", "--output", str(output)],
    )

    assert result.exit_code == 0
    assert called["output"] == output


def _make_db(tmp_path: Path):
    db_path = tmp_path / "jobs.db"
    conn = database.init_db(db_path)
    return db_path, conn


def _insert_job(conn, **overrides) -> None:
    data = {
        "url": "https://example.com/job",
        "title": "Example Job",
        "salary": None,
        "description": "desc",
        "location": "Remote",
        "site": "example",
        "strategy": "manual",
        "discovered_at": "2026-03-10T00:00:00+00:00",
        "full_description": "full desc",
        "application_url": "https://example.com/apply",
        "detail_scraped_at": "2026-03-10T00:01:00+00:00",
        "detail_error": None,
        "fit_score": 9,
        "score_reasoning": "good fit",
        "scored_at": "2026-03-10T00:02:00+00:00",
        "tailored_resume_path": "C:/tmp/resume.txt",
        "tailored_at": "2026-03-10T00:03:00+00:00",
        "tailor_attempts": 0,
        "cover_letter_path": None,
        "cover_letter_at": None,
        "cover_attempts": 0,
        "applied_at": None,
        "apply_status": None,
        "apply_error": None,
        "apply_attempts": 0,
        "agent_id": None,
        "last_attempted_at": None,
        "apply_duration_ms": None,
        "apply_task_id": None,
        "verification_confidence": None,
    }
    data.update(overrides)
    columns = ", ".join(data.keys())
    placeholders = ", ".join("?" for _ in data)
    conn.execute(f"INSERT INTO jobs ({columns}) VALUES ({placeholders})", tuple(data.values()))
    conn.commit()


def _sample_row() -> dict:
    return {
        "url": "https://example.com/job",
        "title": "Example Job",
        "salary": None,
        "description": "desc",
        "location": "Remote",
        "site": "example",
        "strategy": "manual",
        "discovered_at": "2026-03-10T00:00:00+00:00",
        "full_description": "full desc",
        "application_url": "https://example.com/apply",
        "detail_scraped_at": "2026-03-10T00:01:00+00:00",
        "detail_error": None,
        "fit_score": 9,
        "score_reasoning": "good fit",
        "scored_at": "2026-03-10T00:02:00+00:00",
        "tailored_resume_path": "C:/tmp/resume.txt",
        "tailored_at": "2026-03-10T00:03:00+00:00",
        "tailor_attempts": 0,
        "cover_letter_path": None,
        "cover_letter_at": None,
        "cover_attempts": 0,
        "applied_at": None,
        "apply_status": None,
        "apply_error": None,
        "apply_attempts": 0,
        "agent_id": None,
        "last_attempted_at": None,
        "apply_duration_ms": None,
        "apply_task_id": None,
        "verification_confidence": None,
    }


def test_ready_jobs_include_tailored_resume_without_cover_letter(tmp_path) -> None:
    import applypilot.export as export_mod

    db_path, conn = _make_db(tmp_path)
    _insert_job(conn, url="https://example.com/ready-no-cover", cover_letter_path=None)

    rows = export_mod.fetch_ready_jobs_for_export(db_path=db_path)

    assert len(rows) == 1
    assert rows[0]["url"] == "https://example.com/ready-no-cover"
    assert rows[0]["cover_letter_path"] is None


def test_ready_jobs_exclude_applied_rows(tmp_path) -> None:
    import applypilot.export as export_mod

    db_path, conn = _make_db(tmp_path)
    _insert_job(
        conn,
        url="https://example.com/applied",
        applied_at="2026-03-10T00:04:00+00:00",
        apply_status="applied",
    )

    rows = export_mod.fetch_ready_jobs_for_export(db_path=db_path)

    assert rows == []


def test_ready_jobs_exclude_failed_by_default(tmp_path) -> None:
    import applypilot.export as export_mod

    db_path, conn = _make_db(tmp_path)
    _insert_job(
        conn,
        url="https://example.com/failed",
        apply_status="failed",
        apply_error="captcha",
    )

    rows = export_mod.fetch_ready_jobs_for_export(db_path=db_path)

    assert rows == []


def test_ready_jobs_include_failed_when_requested(tmp_path) -> None:
    import applypilot.export as export_mod

    db_path, conn = _make_db(tmp_path)
    _insert_job(
        conn,
        url="https://example.com/failed-include",
        apply_status="failed",
        apply_error="captcha",
    )

    rows = export_mod.fetch_ready_jobs_for_export(db_path=db_path, include_failed=True)

    assert len(rows) == 1
    assert rows[0]["url"] == "https://example.com/failed-include"


def test_export_writes_two_required_sheets(tmp_path) -> None:
    import applypilot.export as export_mod

    output = tmp_path / "ready.xlsx"
    export_mod.export_ready_jobs_to_xlsx(rows=[_sample_row()], output=output)

    workbook = load_workbook(output)

    assert workbook.sheetnames == ["ready_to_apply", "raw_ready_jobs"]


def test_curated_sheet_contains_expected_headers(tmp_path) -> None:
    import applypilot.export as export_mod

    output = tmp_path / "ready.xlsx"
    export_mod.export_ready_jobs_to_xlsx(rows=[_sample_row()], output=output)

    workbook = load_workbook(output)
    worksheet = workbook["ready_to_apply"]
    headers = [cell.value for cell in worksheet[1]]

    assert "tailored_resume_path" in headers
    assert "cover_letter_path" in headers
    assert "application_url" in headers


def test_raw_sheet_contains_full_row_columns(tmp_path) -> None:
    import applypilot.export as export_mod

    output = tmp_path / "ready.xlsx"
    export_mod.export_ready_jobs_to_xlsx(rows=[_sample_row()], output=output)

    workbook = load_workbook(output)
    worksheet = workbook["raw_ready_jobs"]
    headers = [cell.value for cell in worksheet[1]]

    assert "url" in headers
    assert "apply_status" in headers
    assert "score_reasoning" in headers


def test_export_uses_timestamped_default_output(tmp_path) -> None:
    import applypilot.export as export_mod

    path = export_mod.build_default_ready_jobs_export_path(base_dir=tmp_path)

    assert path.parent == tmp_path
    assert path.name.startswith("ready_jobs_")
    assert path.suffix == ".xlsx"


def test_empty_export_still_writes_headers(tmp_path) -> None:
    import applypilot.export as export_mod

    output = tmp_path / "ready.xlsx"
    export_mod.export_ready_jobs_to_xlsx(rows=[], output=output)

    workbook = load_workbook(output)

    assert workbook["ready_to_apply"].max_row == 1
    assert workbook["raw_ready_jobs"].max_row == 1


def test_export_ready_jobs_command_writes_workbook(monkeypatch, tmp_path) -> None:
    import applypilot.export as export_mod

    db_path, conn = _make_db(tmp_path)
    _insert_job(conn, url="https://example.com/export-me", cover_letter_path=None)

    monkeypatch.setattr(cli, "_bootstrap", lambda: None)
    monkeypatch.setattr(database, "DB_PATH", db_path)
    monkeypatch.setattr(export_mod, "APP_DIR", tmp_path)

    output = tmp_path / "manual-ready.xlsx"
    result = runner.invoke(
        cli.app,
        ["export", "ready-jobs", "--output", str(output)],
    )

    assert result.exit_code == 0
    assert output.exists()

    workbook = load_workbook(output)
    assert workbook["ready_to_apply"].max_row == 2


def test_help_command_shows_export_usage() -> None:
    result = runner.invoke(cli.app, ["help", "export"])

    assert result.exit_code == 0
    assert "ready-jobs" in result.output
    assert "xlsx" in result.output.lower()
